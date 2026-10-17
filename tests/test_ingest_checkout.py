from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest

from src.ingest.orbe import ler_orbe
from src.ingest.taxas import atualizar_cadastro_yaml, gerar_cadastro, interpretar_vigencia, ler_taxas

D = Decimal
CAB = "id_transacao;data_movimento;cod_estabelecimento;nome_estabelecimento;tipo_lancamento;modalidade;qtd_parcelas;bandeira;valor_bruto;taxa_aplicada_pct;valor_taxa;valor_liquido;id_referencia;data_venda_original\r\n"
LIN = "TX1;01/09/2026;0123456789;NOME;VENDA;CREDITO PARCELADO;11;ELO;5.773,57;3,79;218,81;5554,76;;\r\n"


def orbe(tmp_path, nome, corpo):
    (tmp_path / nome).write_text(corpo, encoding="utf-8", newline="")


def test_orbe_tipos_e_arquivo_relativo(tmp_path):
    orbe(tmp_path, "01_orbe_x.csv", "linha de titulo\r\n" + CAB + LIN)  # cabeçalho achado pelo conteúdo
    t = ler_orbe(tmp_path)[0]
    assert t["ec"] == "0123456789" and t["valor_bruto"] == D("5773.57")
    assert t["taxa_aplicada_pct"] == D("3.79") and t["qtd_parcelas"] == 11
    assert t["data_movimento"] == "2026-09-01" and t["arquivo_origem"] == "01_orbe_x.csv" and t["ref_origem"] == "TX1"
    assert t["id_referencia"] is None


def test_orbe_agrega_ajustes_e_aceita_ausencia(tmp_path):
    orbe(tmp_path, "01_orbe_x.csv", CAB + LIN)
    assert len(ler_orbe(tmp_path)) == 1
    orbe(tmp_path, "06_orbe_ajustes_x.csv", CAB + LIN.replace("TX1", "TX2"))
    assert [t["arquivo_origem"] for t in ler_orbe(tmp_path)] == ["01_orbe_x.csv", "06_orbe_ajustes_x.csv"]


def test_orbe_id_repetido_e_erro_e_principal_e_obrigatorio(tmp_path):
    with pytest.raises(FileNotFoundError):
        ler_orbe(tmp_path)
    orbe(tmp_path, "01_orbe_x.csv", CAB + LIN)
    orbe(tmp_path, "06_orbe_ajustes_x.csv", CAB + LIN)
    with pytest.raises(ValueError, match="repetido"):
        ler_orbe(tmp_path)


def test_orbe_linha_invalida_diz_arquivo_e_linha(tmp_path):
    orbe(tmp_path, "01_orbe_x.csv", CAB + LIN.replace("5.773,57", "abc"))
    with pytest.raises(ValueError, match=r"01_orbe_x.csv, linha 2"):
        ler_orbe(tmp_path)


@pytest.mark.parametrize("texto,esperado", [
    (None, (None, None)), ("até 14/09/2026", (None, "2026-09-14")),
    ("a partir de 15/09/2026", ("2026-09-15", None)), ("01/01/2026", ("2026-01-01", None)),
    ("01/01/2026 a 30/06/2026", ("2026-01-01", "2026-06-30")),
])
def test_vigencia(texto, esperado):
    assert interpretar_vigencia(texto) == esperado


def test_vigencia_desconhecida_e_erro():
    with pytest.raises(ValueError):
        interpretar_vigencia("quando der")


def planilha(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Taxas vigentes"
    ws["A1"] = "título"
    ws.append([])
    cab = ["Lojista", "CNPJ", "EC Orbe", "Débito", "Crédito à vista", "Créd. 2-6x", "Créd. 7-12x", "Pix", "Antecipação", "Spread remessa", "Vigência", "Obs"]
    ws.append(cab)
    ws.append(["Loja A", "11.222.333/0001-44", "1000000001", 0.0149, "2,59%", 0.0319, 0.0379, 0.0089, 0.021, None, "01/01/2026", None])
    ws.append(["Loja B", 3229837000103, 1000000002, 0.0149, 0.0269, 0.0319, 0.0379, 0.0089, 0.021, "padrão", None, "INATIVO desde jul/26"])
    ws.append(["Só remessa", "5.000.000/0001-00", None, None, None, None, None, None, None, 0.0125, "01/01/2026", "só remessa"])
    ws.append([])
    ws.append(["* spread em branco = padrão"])
    velha = wb.create_sheet("ANTIGO 2025 - nao usar")
    velha.append(["Lojista", "CNPJ"])
    velha.append(["Fantasma", "99999999000199"])
    p = tmp_path / "taxas.xlsx"
    wb.save(p)
    return p


def test_taxas_aba_vigente_tipos_mistos_e_refs(tmp_path):
    cs = ler_taxas(planilha(tmp_path))
    assert [c["lojista"] for c in cs] == ["Loja A", "Loja B", "Só remessa"]  # sem rodapé nem aba antiga
    a, b, r = cs
    assert a["cnpj"] == "11222333000144" and a["ref_origem"] == "Taxas vigentes!L4"
    assert a["taxas_pct"]["CREDITO_A_VISTA"] == D("2.59") == D("2.59")  # '2,59%'
    assert a["taxas_pct"]["DEBITO"] == D("1.49")  # 0.0149 sem ruído de float
    assert b["cnpj"] == "03229837000103" and b["ec"] == "1000000002" and b["ativo"] is False
    assert (b["vigencia_inicio"], b["vigencia_fim"]) == (None, None)
    assert all(v is None for v in r["taxas_pct"].values()) and r["ec"] is None
    # Alvo: antecipação e spread de remessa; vazio e "padrão" = None (spread padrão, P4)
    assert a["taxas_pct"]["ANTECIPACAO"] == D("2.10")
    assert a["spread_remessa_pct"] is None and b["spread_remessa_pct"] is None
    assert r["spread_remessa_pct"] == D("1.25")


def test_cadastro_so_com_ec_e_nao_sobrescreve(tmp_path):
    cs = ler_taxas(planilha(tmp_path))
    assert set(gerar_cadastro(cs)) == {"1000000001", "1000000002"}
    destino = tmp_path / "lojistas.yaml"
    assert atualizar_cadastro_yaml(cs, destino) == 2
    assert atualizar_cadastro_yaml(cs, destino) == 0  # idempotente
    import yaml
    assert yaml.safe_load(destino.read_text(encoding="utf-8"))["lojistas"]["1000000002"]["cnpj"] == "03229837000103"


def test_cadastro_ec_com_dois_cnpjs_e_erro(tmp_path):
    cs = ler_taxas(planilha(tmp_path))
    cs[1] = {**cs[1], "ec": cs[0]["ec"]}
    with pytest.raises(ValueError, match="CNPJs diferentes"):
        gerar_cadastro(cs)


def test_data_pack_taxas_e_cadastro():
    pack = Path(__file__).resolve().parent.parent / "context" / "tarvo_data_pack"
    if not pack.exists():
        pytest.skip("data pack ausente")
    cs = ler_taxas(pack / "02_tabela_taxas_contratadas.xlsx")
    assert len(cs) == 25 and len(gerar_cadastro(cs)) == 19
    bridge = [c for c in cs if c["ec"] == "1075624862"]
    assert [(c["vigencia_inicio"], c["vigencia_fim"]) for c in bridge] == [(None, "2026-09-14"), ("2026-09-15", None)]
    assert [c["taxas_pct"]["CREDITO_A_VISTA"] for c in bridge] == [D("2.59"), D("2.19")]
    assert len(ler_orbe(pack)) == 1957
    spreads = {c["lojista"]: c["spread_remessa_pct"] for c in cs}
    assert spreads["Mundo Aberto"] == D("1.20") and spreads["Canguru"] is None  # 'padrão'
    assert sum(1 for c in cs if c["taxas_pct"]["ANTECIPACAO"] == D("2.10")) == 23
