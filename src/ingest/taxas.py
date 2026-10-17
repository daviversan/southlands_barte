"""Leitor da planilha de taxas contratadas (só a aba vigente) e gerador do cadastro mestre.

A planilha é mantida à mão: taxas em fração e em texto com `%`, CNPJ como número, vigência em
texto livre. Cabeçalho localizado pelo conteúdo; a aba `ANTIGO 2025 - nao usar` é ignorada.
Uma linha por vigência. Taxas em pontos percentuais; ausente = None (nunca zero).
"""
from __future__ import annotations

import re
import unicodedata
from decimal import Decimal
from pathlib import Path

import openpyxl
import yaml

from src.core import normalize
from src.core.contratos import ContratoTaxa

ABA = "taxas vigentes"

# cabeçalho normalizado -> campo interno
_COLUNAS = {
    "lojista": "lojista",
    "cnpj": "cnpj",
    "ec orbe": "ec",
    "debito": "DEBITO",
    "credito a vista": "CREDITO_A_VISTA",
    "cred. 2-6x": "CREDITO_2_6",
    "cred. 7-12x": "CREDITO_7_12",
    "pix": "PIX",
    "antecipacao": "ANTECIPACAO",
    "spread remessa": "spread_remessa",
    "vigencia": "vigencia",
    "obs": "obs",
}
_FAIXAS = ("DEBITO", "CREDITO_A_VISTA", "CREDITO_2_6", "CREDITO_7_12", "PIX", "ANTECIPACAO")
_SPREAD_PADRAO = ("padrao",)  # texto que significa "spread padrão" (P4 aberta): vira None
_DATA = re.compile(r"\d{2}/\d{2}/\d{4}")


def _chave(texto) -> str:
    t = unicodedata.normalize("NFKD", str(texto)).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", t).strip().lower()


def _achar_aba(wb: openpyxl.Workbook):
    for ws in wb.worksheets:
        if _chave(ws.title) == ABA:
            return ws
    raise ValueError(f"aba '{ABA}' não encontrada; abas: {wb.sheetnames}")


def _achar_cabecalho(ws) -> tuple[int, dict[str, int]]:
    """(linha física, {campo: índice da coluna}) da primeira linha que tem Lojista e CNPJ."""
    for numero, linha in enumerate(ws.iter_rows(values_only=True), start=1):
        mapa = {_COLUNAS[_chave(v)]: i for i, v in enumerate(linha)
                if v is not None and _chave(v) in _COLUNAS}
        if "lojista" in mapa and "cnpj" in mapa:
            return numero, mapa
    raise ValueError("cabeçalho (Lojista, CNPJ) não encontrado na aba de taxas vigentes")


def _texto(x) -> str | None:
    return None if normalize._vazio(x) else str(x).strip()


def _ec(x) -> str | None:
    if normalize._vazio(x):
        return None
    if isinstance(x, float) and x == int(x):
        x = int(x)
    return str(x).strip()


def _spread(x) -> Decimal | None:
    """Spread de remessa em pontos percentuais. Vazio ou `padrão` = None (spread padrão, P4)."""
    if normalize._vazio(x) or (isinstance(x, str) and _chave(x) in _SPREAD_PADRAO):
        return None
    return normalize.pct(x, fracao=True)


def interpretar_vigencia(x) -> tuple[str | None, str | None]:
    """Texto de vigência -> (início, fim) ISO, ambos inclusivos; None = limite aberto."""
    if normalize._vazio(x):
        return None, None
    if not isinstance(x, str):
        return normalize.data(x), None  # data solta (célula de data): início
    datas = [normalize.data(d) for d in _DATA.findall(x)]
    t = _chave(x)
    if len(datas) == 2:
        return datas[0], datas[1]
    if len(datas) == 1:
        if t.startswith("ate"):
            return None, datas[0]
        return datas[0], None  # "a partir de", "desde" ou data solta
    raise ValueError(f"vigência não reconhecida: {x!r}")


def ler_taxas(caminho: Path, arquivo_origem: str | None = None) -> list[ContratoTaxa]:
    """Contratos da aba vigente, um por linha (vigência). `arquivo_origem` default = nome do arquivo."""
    caminho = Path(caminho)
    rotulo = arquivo_origem or caminho.name
    wb = openpyxl.load_workbook(caminho, data_only=True, read_only=False)
    ws = _achar_aba(wb)
    cab, col = _achar_cabecalho(ws)

    def celula(linha, campo):
        i = col.get(campo)
        return linha[i] if i is not None and i < len(linha) else None

    contratos: list[ContratoTaxa] = []
    for numero, linha in enumerate(ws.iter_rows(min_row=cab + 1, values_only=True), start=cab + 1):
        lojista = _texto(celula(linha, "lojista"))
        if lojista is None or lojista.startswith("*"):  # linha vazia ou nota de rodapé
            continue
        ref = f"{ws.title}!L{numero}"
        try:
            cnpj = normalize.cnpj(celula(linha, "cnpj"))
            if cnpj is None:
                raise ValueError("CNPJ vazio")
            ini, fim = interpretar_vigencia(celula(linha, "vigencia"))
            obs = _chave(celula(linha, "obs") or "")
            contratos.append({
                "cnpj": cnpj,
                "lojista": lojista,
                "ec": _ec(celula(linha, "ec")),
                "ativo": "inativ" not in obs,
                "vigencia_inicio": ini,
                "vigencia_fim": fim,
                "taxas_pct": {f: normalize.pct(celula(linha, f), fracao=True) for f in _FAIXAS},
                "spread_remessa_pct": _spread(celula(linha, "spread_remessa")),
                "arquivo_origem": rotulo,
                "ref_origem": ref,
            })
        except ValueError as e:
            raise ValueError(f"{rotulo}, {ref}: {e}") from e
    return contratos


# --- cadastro mestre (EC -> CNPJ) -------------------------------------------------------------

_CABECALHO_YAML = """\
# Cadastro mestre de lojistas (EC <-> CNPJ).
# Gerado por regra EXATA (EC da planilha de taxas) e completado só com confirmação humana.
# Nunca por fuzzy match de nome. Regerar: python -m src.ingest.taxas <planilha.xlsx> config/lojistas.yaml
# A regeração só ACRESCENTA ECs novos; entradas existentes (inclusive confirmadas à mão) são mantidas.
"""


class _DumperAspas(yaml.SafeDumper):
    """Todo texto entre aspas: CNPJ/EC com zero à esquerda nunca viram número em outro leitor de YAML."""


_DumperAspas.add_representer(str, lambda d, v: d.represent_scalar("tag:yaml.org,2002:str", v, style='"'))


def gerar_cadastro(contratos: list[ContratoTaxa]) -> dict[str, dict[str, str]]:
    """EC -> {cnpj, lojista}, só de contratos com EC. Mesmo EC com CNPJ diferente é erro."""
    cadastro: dict[str, dict[str, str]] = {}
    for c in contratos:
        if c["ec"] is None:
            continue
        atual = cadastro.setdefault(c["ec"], {"cnpj": c["cnpj"], "lojista": c["lojista"]})
        if atual["cnpj"] != c["cnpj"]:
            raise ValueError(f"EC {c['ec']} com CNPJs diferentes: {atual['cnpj']} e {c['cnpj']} ({c['ref_origem']})")
    return cadastro


def atualizar_cadastro_yaml(contratos: list[ContratoTaxa], destino: Path) -> int:
    """Acrescenta ao YAML os ECs ainda ausentes. Devolve quantos entraram; conflito de CNPJ é erro."""
    destino = Path(destino)
    existente = {}
    if destino.exists():
        existente = (yaml.safe_load(destino.read_text(encoding="utf-8")) or {}).get("lojistas") or {}
    novos = gerar_cadastro(contratos)
    for ec, dados in novos.items():
        if ec in existente and existente[ec]["cnpj"] != dados["cnpj"]:
            raise ValueError(f"EC {ec}: cadastro tem {existente[ec]['cnpj']}, planilha tem {dados['cnpj']}")
    acrescentar = {ec: d for ec, d in novos.items() if ec not in existente}
    final = {**existente, **acrescentar}
    corpo = yaml.dump({"lojistas": final}, Dumper=_DumperAspas, allow_unicode=True, sort_keys=True, default_flow_style=False)
    destino.write_text(_CABECALHO_YAML + corpo, encoding="utf-8")
    return len(acrescentar)


if __name__ == "__main__":
    import sys

    if len(sys.argv) != 3:
        sys.exit("uso: python -m src.ingest.taxas <planilha.xlsx> <lojistas.yaml>")
    n = atualizar_cadastro_yaml(ler_taxas(Path(sys.argv[1])), Path(sys.argv[2]))
    print(f"{n} lojista(s) acrescentado(s) em {sys.argv[2]}")
